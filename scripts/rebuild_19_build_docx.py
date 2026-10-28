# -*- coding: utf-8 -*-
"""
Rebuild submission/manuscript_final_with_figures.docx from the HONEST
manuscript_final.md.

- Reads the real title from the markdown (no hardcoded old title).
- Embeds the 12 genuine figures from results/ (incl. regenerated Fig 2 and
  the new Fig 10 structural / Fig 11 audit figures).
- Retired MR (Fig 10 old) / MD (Fig 11 old) / Fig 12 are NOT embedded.
- Supplementary Materials section is rendered with headings/tables (not forced
  into a figure block).
"""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = r"D:/省自然/生信分析"
MD = os.path.join(BASE, "submission", "manuscript_final.md")
OUT = os.path.join(BASE, "submission", "manuscript_final_with_figures.docx")
FIGDIR = os.path.join(BASE, "results")

# figure number -> filename in results/
FIG_FILES = {
    "Figure 0":  "Figure0_Graphical_Abstract.png",
    "Figure 1":  "Figure1_DEG_overview.png",
    "Figure 2":  "Figure2_Wnt_pathway.png",
    "Figure 3":  "Figure3_Pyroptosis_Wnt_crosstalk.png",
    "Figure 4":  "Figure4_scRNAseq_cell_atlas.png",
    "Figure 5":  "Figure5_CellChat_Pseudotime.png",
    "Figure 6":  "Figure6_Deconvolution.png",
    "Figure 7":  "Figure7_PPI_Hub_TCF4.png",
    "Figure 8":  "Figure8_Drug_ThreeTier_Classification.png",
    "Figure 9":  "Figure9_GRN_VirtualKO.png",
    "Figure 10": "Figure10_Structural_Analysis.png",
    "Figure 11": "Figure11_Reproducibility_Audit.png",
}

with open(MD, encoding="utf-8") as fh:
    raw = fh.read()

legend_marker = "## FIGURE LEGENDS"
idx = raw.find(legend_marker)
head_text = raw[:idx].rstrip()
legend_text = raw[idx:].rstrip()

# --- extract title (first '# ' line) ---
hlines = head_text.split("\n")
title = None
start = 0
for k, ln in enumerate(hlines):
    if ln.strip().startswith("# "):
        title = ln.strip()[2:].strip()
        start = k + 1
        break
body_text = "\n".join(hlines[start:]).strip()


def _add_runs_with_bold(p, text):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)


def _add_table(doc, tbl_lines):
    rows = []
    for ln in tbl_lines:
        inner = ln.strip().strip("|")
        cells = [c.strip() for c in inner.split("|")]
        rows.append(cells)
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Light Grid Accent 1"
    for cells in rows:
        row_cells = table.add_row().cells
        for j in range(n_cols):
            row_cells[j].text = (cells[j] if j < len(cells) else "").replace("**", "")
    doc.add_paragraph()


def add_markdown(doc, text):
    text = text.replace("\r\n", "\n")
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(m.group(2).strip().replace("**", ""), level=min(level, 3))
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
            i += 1; continue
        if stripped.startswith("---"):
            i += 1; continue
        if stripped.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip()); i += 1
            _add_table(doc, tbl); continue
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(stripped[2:].strip().replace("**", ""), style="List Bullet")
            p.paragraph_format.space_after = Pt(2); i += 1; continue
        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(re.sub(r"^\d+\.\s+", "", stripped).replace("**", ""), style="List Number")
            p.paragraph_format.space_after = Pt(2); i += 1; continue
        if stripped.startswith("**") and stripped.endswith("**"):
            p = doc.add_paragraph()
            r = p.add_run(stripped.strip("*")); r.bold = True; r.font.size = Pt(11)
            i += 1; continue
        p = doc.add_paragraph()
        _add_runs_with_bold(p, stripped)
        p.paragraph_format.space_after = Pt(6)
        i += 1


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"; style.font.size = Pt(11)

# Title (from md)
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = tp.add_run(title); tr.bold = True; tr.font.size = Pt(15)
doc.add_paragraph()

add_markdown(doc, body_text)

# --- FIGURE LEGENDS + embedded figures ---
doc.add_page_break()
doc.add_heading("FIGURE LEGENDS AND FIGURES", level=1)

blocks = re.split(r"(?=^\*\*Figure \d+ –)", legend_text, flags=re.M)
embedded = 0
for blk in blocks:
    blk = blk.strip()
    if not blk:
        continue
    m = re.match(r"^\*\*(Figure \d+)\s*[–-]?\s*(.*)$", blk, flags=re.S)
    if m:
        fig_key = m.group(1)
        rest = m.group(2)
        rest = rest.replace(".** ", ". ", 1)  # strip spurious bold closer after title
        p = doc.add_paragraph()
        r = p.add_run(f"{fig_key}  "); r.bold = True; r.font.size = Pt(11.5)
        _add_runs_with_bold(p, rest)
        if fig_key in FIG_FILES:
            img = os.path.join(FIGDIR, FIG_FILES[fig_key])
            if os.path.exists(img):
                pict = doc.add_paragraph(); pict.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pict.add_run().add_picture(img, width=Inches(6.3))
                embedded += 1
            else:
                doc.add_paragraph(f"[IMAGE MISSING: {img}]")
    else:
        # supplementary materials / other trailing text -> render as markdown
        add_markdown(doc, blk)

doc.save(OUT)
print(f"Saved: {OUT}")
print(f"Figures embedded: {embedded} / {len(FIG_FILES)}")
